import io
import json
import pytest
import requests
import openpyxl
import pypdf

BASE_URL = "http://127.0.0.1:8000"
OLLAMA_URL = "http://localhost:11434"


def get_auth_token():
    """Obtiene un token JWT real desde el endpoint de autenticación de FastAPI."""
    res = requests.post(
        f"{BASE_URL}/api/v1/auth/login",
        json={"email": "admin@epidemic-twin.ke", "password": "AdminSecure2026!"},
        timeout=10
    )
    assert res.status_code == 200, f"Login falló: {res.status_code} {res.text}"
    data = res.json()
    token = data.get("access_token")
    assert token, "No se recibió access_token en el login"
    return token


# ===========================================================================
# 1. Pruebas que fallan si los servicios reales no responden
# ===========================================================================

def test_live_fastapi_login_responds():
    """Falla si el backend FastAPI o el login real no responden."""
    token = get_auth_token()
    assert len(token) > 20, "Token JWT inválido o vacío"


def test_live_ollama_service_and_model():
    """Falla si Ollama no está corriendo en localhost:11434 o falta el modelo qwen2.5:0.5b."""
    # 1. Verificar tags
    tags_res = requests.get(f"{OLLAMA_URL}/api/tags", timeout=5)
    assert tags_res.status_code == 200, f"Ollama /api/tags no respondió 200: {tags_res.status_code}"
    tags_data = tags_res.json()
    models = [m.get("name") for m in tags_data.get("models", [])]
    assert any("qwen2.5:0.5b" in m for m in models), f"Modelo qwen2.5:0.5b no encontrado en Ollama: {models}"

    # 2. Generación real
    gen_res = requests.post(
        f"{OLLAMA_URL}/api/generate",
        json={"model": "qwen2.5:0.5b", "prompt": "Responder con una palabra: OK", "stream": False},
        timeout=30
    )
    assert gen_res.status_code == 200, f"Ollama /api/generate falló: {gen_res.status_code}"
    response_text = gen_res.json().get("response", "")
    assert len(response_text.strip()) > 0, "Ollama devolvió respuesta vacía"


def test_live_langflow_query_responds():
    """Falla si el contenedor de Langflow en Docker o la consulta en FastAPI no responden."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}
    res = requests.post(
        f"{BASE_URL}/api/v1/langflow/query",
        json={"query": "¿Cuál es el modelo campeón de vigilancia epidemiológica?"},
        headers=headers,
        timeout=45
    )
    assert res.status_code == 200, f"Langflow query falló: {res.status_code} {res.text}"
    data = res.json()
    assert data.get("status", "").upper() == "SUCCESS", f"Estado no exitoso: {data}"
    reply = data.get("reply", "")
    assert len(reply) > 10, f"Respuesta de Langflow insuficiente: {reply}"
    # Verificar que el modelo usado o la respuesta contenga datos verificables
    assert "Naive" in reply or "campeón" in reply or "Vigilancia" in reply or "Baseline" in reply or "CV" in reply, f"Respuesta no esperada: {reply}"


def test_live_chatbot_ollama_responds():
    """Falla si el Copiloto IA (/api/v1/ai/chat) no responde usando Ollama."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}
    res = requests.post(
        f"{BASE_URL}/api/v1/ai/chat",
        json={"message": "Responde en 1 frase: ¿Qué es la vigilancia epidemiológica semanal?", "conversation_history": []},
        headers=headers,
        timeout=60
    )
    assert res.status_code == 200, f"Chatbot falló: {res.status_code} {res.text}"
    data = res.json()
    model = data.get("model", "")
    reply = data.get("reply", "")
    assert "Ollama" in model or "qwen2.5" in model, f"Modelo no corresponde a Ollama: {model}"
    assert len(reply) > 5, "Respuesta vacía del copiloto"


def test_live_academic_translator_ollama():
    """Falla si el Academic Translator no realiza traducción mediante Ollama."""
    from apps.academic_translator.agents.translator_pipeline import AcademicTranslatorPipeline
    pipeline = AcademicTranslatorPipeline()
    health = pipeline.check_health()
    assert health is True, f"Health check de traductor falló con base_url: {pipeline.base_url}"

    result = pipeline.translate_document(
        text="Weekly incidence decreased following non-pharmaceutical interventions.",
        filename="test_paper.md",
        target_lang="es"
    )
    assert result.get("translated_text"), "Texto traducido vacío"
    assert "ollama" in result.get("model_used", "").lower(), f"Modelo no es Ollama: {result.get('model_used')}"


# ===========================================================================
# 2. Pruebas de Cálculos Dinámicos y Reportes sin Valores Escritos a Mano
# ===========================================================================

def test_reports_different_cutoffs_produce_different_predictions():
    """Verifica que una fecha de corte diferente produzca un reporte de predicción diferente."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    # Corte 1: 2024-07-28 (cierre endémico, casos t0 = 0)
    res_2024 = requests.post(
        f"{BASE_URL}/api/v1/reports/preview",
        json={"report_type": "prediction", "cutoff_date": "2024-07-28", "lang": "es"},
        headers=headers,
        timeout=10
    )
    assert res_2024.status_code == 200
    data_2024 = res_2024.json()

    # Corte 2: 2021-07-11 (pico ola Delta, casos t0 = 2922)
    res_2021 = requests.post(
        f"{BASE_URL}/api/v1/reports/preview",
        json={"report_type": "prediction", "cutoff_date": "2021-07-11", "lang": "es"},
        headers=headers,
        timeout=10
    )
    assert res_2021.status_code == 200
    data_2021 = res_2021.json()

    # Extraer valores de métricas
    val_2024 = next(m["value"] for m in data_2024["summary_metrics"] if "Pronóstico Oficial" in m["label"])
    val_2021 = next(m["value"] for m in data_2021["summary_metrics"] if "Pronóstico Oficial" in m["label"])

    assert val_2024 != val_2021, f"Ambos cortes produjeron el mismo valor: {val_2024}"
    assert "0.00" in val_2024, f"Se esperaba 0.00 en 2024-07-28 pero fue: {val_2024}"
    assert "2922.00" in val_2021, f"Se esperaba 2922.00 en 2021-07-11 pero fue: {val_2021}"


def test_reports_metrics_match_json_including_ridge_mae():
    """Verifica que las métricas del reporte coinciden con model_evaluation_weekly_results.json."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    res = requests.post(
        f"{BASE_URL}/api/v1/reports/preview",
        json={"report_type": "models", "lang": "es"},
        headers=headers,
        timeout=10
    )
    assert res.status_code == 200
    data = res.json()

    # Tabla 1 de comparación de modelos
    table1 = data["tables"][0]
    models_rows = {row[0]: row for row in table1["rows"]}

    # Ridge MAE debe ser exactamente 16.04 (no 16.78 escrito a mano)
    ridge_row = models_rows.get("Baseline_Ridge_Regression")
    assert ridge_row is not None, "Baseline_Ridge_Regression no encontrada en tabla"
    # Columnas: Modelo, Rol, CV RMSE (Dev), Test RMSE, Test MAE, Test MAPE
    ridge_test_mae = ridge_row[4]
    ridge_cv_rmse = ridge_row[2]
    assert ridge_test_mae == "16.04", f"Ridge Test MAE esperado '16.04', obtenido: '{ridge_test_mae}'"
    assert ridge_cv_rmse == "1,405.93", f"Ridge CV RMSE esperado '1,405.93', obtenido: '{ridge_cv_rmse}'"

    # Naive Persistence CV RMSE debe ser 1183.64
    naive_row = models_rows.get("Baseline_Naive_Persistence")
    assert naive_row is not None
    assert naive_row[2] == "1,183.64"


def test_reports_scenarios_reads_persisted_simulation():
    """Verifica que el reporte de escenarios lee la simulación persistida real."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    res = requests.post(
        f"{BASE_URL}/api/v1/reports/preview",
        json={"report_type": "scenarios", "lang": "es"},
        headers=headers,
        timeout=10
    )
    assert res.status_code == 200
    data = res.json()
    metrics = {m["label"]: m["value"] for m in data["summary_metrics"]}
    assert "ID de Simulación Persistida" in metrics
    sim_id = metrics["ID de Simulación Persistida"]
    assert sim_id.startswith("sim-"), f"ID de simulación inválido: {sim_id}"
    assert "Tasa de Ataque Proyectada" in metrics


# ===========================================================================
# 3. Exportación, Apertura e Inspección de Binarios (PDF y Excel)
# ===========================================================================

def test_export_and_inspect_pdf_contains_selected_date():
    """Exporta un PDF real, lo abre con pypdf y comprueba que contiene la fecha elegida."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}
    cutoff = "2021-07-11"

    res = requests.post(
        f"{BASE_URL}/api/v1/reports/export/pdf",
        json={"report_type": "prediction", "cutoff_date": cutoff, "lang": "es"},
        headers=headers,
        timeout=15
    )
    assert res.status_code == 200, f"Export PDF falló: {res.status_code}"
    pdf_bytes = res.content
    assert len(pdf_bytes) > 1000, "PDF generado demasiado pequeño"

    # Abrir el PDF e inspeccionar su texto con pypdf
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1, "PDF no contiene páginas"

    extracted_text = ""
    for page in reader.pages:
        extracted_text += page.extract_text() or ""

    # Comprobar que contiene la fecha elegida y la cifra correspondiente
    assert cutoff in extracted_text, f"Fecha de corte {cutoff} no encontrada en el PDF"
    assert "2922.00" in extracted_text or "2922" in extracted_text, "Cifra observada 2922 no encontrada en el PDF"
    assert "GD_TI Digital Twin Kenya" in extracted_text


def test_export_and_inspect_excel_contains_selected_simulation():
    """Exporta un Excel real, lo abre con openpyxl y comprueba que contiene la simulación elegida."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    res = requests.post(
        f"{BASE_URL}/api/v1/reports/export/excel",
        json={"report_type": "scenarios", "lang": "en"},
        headers=headers,
        timeout=15
    )
    assert res.status_code == 200, f"Export Excel falló: {res.status_code}"
    excel_bytes = res.content
    assert len(excel_bytes) > 1000, "Excel generado demasiado pequeño"

    # Abrir el Excel e inspeccionar celdas con openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert len(sheet_names) >= 2, f"Se esperaban al menos 2 hojas en Excel: {sheet_names}"

    # Hoja de resumen (en inglés: Fact Sheet & Traceability)
    ws1 = wb[sheet_names[0]]
    all_values = []
    for row in ws1.iter_rows(values_only=True):
        for cell in row:
            if cell is not None:
                all_values.append(str(cell))

    joined_text = " ".join(all_values)
    assert "sim-" in joined_text, "No se encontró ID de simulación en las celdas de Excel"
    assert "SEIR/ABM" in joined_text or "Simulation" in joined_text or "Attack Rate" in joined_text


# ===========================================================================
# 4. Pruebas Separadas para Verificar Mensajes de Error Cuando Están Apagados
# ===========================================================================

def test_offline_error_handling_when_ollama_unreachable():
    """Verifica que el traductor académico maneja ordenadamente cuando Ollama está apagado."""
    from apps.academic_translator.agents.translator_pipeline import AcademicTranslatorPipeline
    pipeline = AcademicTranslatorPipeline(base_url="http://127.0.0.1:59999")
    assert pipeline.check_health() is False, "Se esperaba False en check_health cuando el servicio está apagado"
    with pytest.raises(RuntimeError) as exc_info:
        pipeline.translate_document("Texto de prueba", "test.md", "es")
    err_str = str(exc_info.value)
    assert "No se pudo conectar con el servicio local Ollama" in err_str
    assert "ollama serve" in err_str


# ===========================================================================
# 5. Pruebas que Provocan Fallos y Verifican Cero Resultados de Respaldo Ficticios
# ===========================================================================

def test_provoke_failure_invalid_simulation_id_rejects_without_fallback():
    """Provoca un fallo por ID de simulación inexistente y confirma rechazo HTTP 404 sin generar PDF/Excel ficticio."""
    from fastapi.testclient import TestClient
    from backend.app.main import app

    client = TestClient(app)
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}
    bad_id = "sim-invalido-inexistente-99999"

    # 1. Preview debe retornar 404 explícito
    prev_res = client.post(
        "/api/v1/reports/preview",
        json={"report_type": "scenarios", "simulation_id": bad_id},
        headers=headers
    )
    assert prev_res.status_code == 404, f"Se esperaba 404 para simulación inválida, recibido {prev_res.status_code}"
    assert "inexistente en la base de datos" in prev_res.json().get("detail", "")

    # 2. Export PDF debe fallar con 404 y NO devolver un documento binario
    pdf_res = client.post(
        "/api/v1/reports/export/pdf",
        json={"report_type": "scenarios", "simulation_id": bad_id},
        headers=headers
    )
    assert pdf_res.status_code == 404, f"Se esperaba 404 en PDF para simulación inválida, recibido {pdf_res.status_code}"
    assert pdf_res.headers.get("content-type") == "application/json"

    # 3. Export Excel debe fallar con 404 y NO devolver una hoja binaria
    excel_res = client.post(
        "/api/v1/reports/export/excel",
        json={"report_type": "scenarios", "simulation_id": bad_id},
        headers=headers
    )
    assert excel_res.status_code == 404, f"Se esperaba 404 en Excel para simulación inválida, recibido {excel_res.status_code}"
    assert excel_res.headers.get("content-type") == "application/json"


def test_provoke_failure_prediction_inference_rejects_without_zeros():
    """Provoca un fallo de inferencia por fecha inválida y confirma rechazo explícito sin emitir ceros ni archivos."""
    from fastapi.testclient import TestClient
    from backend.app.main import app

    client = TestClient(app)
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    # Fecha no válida
    bad_cutoff = "2099-12-31"

    prev_res = client.post(
        "/api/v1/reports/preview",
        json={"report_type": "prediction", "cutoff_date": bad_cutoff},
        headers=headers
    )
    assert prev_res.status_code == 404, f"Se esperaba 404 para fecha fuera de rango sin entradas, recibido {prev_res.status_code}"
    assert "no encontrada" in prev_res.json().get("detail", "")

    pdf_res = client.post(
        "/api/v1/reports/export/pdf",
        json={"report_type": "prediction", "cutoff_date": bad_cutoff},
        headers=headers
    )
    assert pdf_res.status_code == 404, f"Se esperaba 404 en PDF para fecha fuera de rango, recibido {pdf_res.status_code}"
    assert pdf_res.headers.get("content-type") == "application/json"


def test_provoke_failure_missing_raw_csv_rejects_without_hash():
    """Provoca la ausencia del CSV de procedencia y confirma que no se devuelve el hash de respaldo."""
    from fastapi.testclient import TestClient
    from backend.app.main import app
    from backend.app.api.v1.endpoints import reports

    client = TestClient(app)
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    orig_path = reports.RAW_DATA_PATH
    try:
        reports.RAW_DATA_PATH = "non_existent_fake_path.csv"
        res = client.post(
            "/api/v1/reports/preview",
            json={"report_type": "data_provenance"},
            headers=headers
        )
        assert res.status_code == 404, f"Se esperaba 404 para CSV ausente, recibido {res.status_code}"
        assert "Archivo de datos crudos ausente" in res.json().get("detail", "")
        # Jamás debe retornar el hash estático en caso de ausencia
        assert "c82f3cfc" not in str(res.json())

        # Verificar también que la exportación PDF falla
        pdf_res = client.post(
            "/api/v1/reports/export/pdf",
            json={"report_type": "data_provenance"},
            headers=headers
        )
        assert pdf_res.status_code == 404
    finally:
        reports.RAW_DATA_PATH = orig_path


def test_provoke_failure_missing_eval_json_rejects_without_metrics():
    """Provoca la ausencia del JSON de evaluación y confirma que no se emiten métricas inventadas."""
    from fastapi.testclient import TestClient
    from backend.app.main import app
    from backend.app.api.v1.endpoints import reports

    client = TestClient(app)
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    orig_path = reports.EVAL_RESULTS_PATH
    try:
        reports.EVAL_RESULTS_PATH = "non_existent_eval_results.json"
        res = client.post(
            "/api/v1/reports/preview",
            json={"report_type": "models"},
            headers=headers
        )
        assert res.status_code == 404, f"Se esperaba 404 para JSON ausente, recibido {res.status_code}"
        assert "Archivo de evaluación de modelos ausente" in res.json().get("detail", "")

        # Verificar que Excel no se genera
        xl_res = client.post(
            "/api/v1/reports/export/excel",
            json={"report_type": "models"},
            headers=headers
        )
        assert xl_res.status_code == 404
    finally:
        reports.EVAL_RESULTS_PATH = orig_path


def test_scenarios_report_clearly_declares_theoretical_simulation_condition():
    """Verifica que el reporte de escenarios declara inequívocamente su condición de simulación matemática teórica."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    res = requests.post(
        f"{BASE_URL}/api/v1/reports/preview",
        json={"report_type": "scenarios", "lang": "es"},
        headers=headers,
        timeout=15
    )
    assert res.status_code == 200, f"Error en preview de escenarios: {res.status_code}"
    data = res.json()
    badge = data.get("provenance_badge", "")
    subtitle = data.get("subtitle", "")
    notes = " ".join(data.get("scientific_notes", []))

    assert "SIMULACIÓN MATEMÁTICA TEÓRICA" in badge
    assert "Simulación Matemática Teórica" in subtitle
    assert "no sustituyen datos de vigilancia empírica" in notes or "proyecciones contrafactuales" in notes


def test_prediction_report_alignment_strictly_t_plus_1():
    """Verifica que el reporte de predicción presente estrictamente el horizonte a un paso t+1."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    res = requests.post(
        f"{BASE_URL}/api/v1/reports/preview",
        json={"report_type": "prediction", "cutoff_date": "2024-07-28", "lang": "es"},
        headers=headers,
        timeout=15
    )
    assert res.status_code == 200, f"Error en preview de predicción: {res.status_code}"
    data = res.json()
    title = data.get("title", "")
    subtitle = data.get("subtitle", "")
    notes = " ".join(data.get("scientific_notes", []))

    assert "t+1" in title
    assert "1 paso a t+1" in subtitle
    assert "horizonte de pronóstico es estrictamente a un paso semanal (t+1" in notes


def test_provoke_failure_langchain_agent_ollama_unreachable_rejects_without_fallback():
    """Verifica que si Ollama no está disponible, el endpoint LangChain responde HTTP 503 sin emitir la respuesta determinista."""
    from fastapi.testclient import TestClient
    from backend.app.main import app
    import os

    client = TestClient(app)
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    # Configurar URL de Ollama a un puerto inaccesible
    orig_url = os.environ.get("OLLAMA_BASE_URL", "http://127.0.0.1:11434")
    try:
        os.environ["OLLAMA_BASE_URL"] = "http://127.0.0.1:59999"
        res = client.post(
            "/api/v1/langchain-agent/query",
            json={"query": "¿Cuál es la proyección epidemiológica?", "cutoff_date": "2024-07-28", "lang": "es"},
            headers=headers
        )
        assert res.status_code == 503, f"Se esperaba 503 cuando Ollama es inaccesible, recibido: {res.status_code}"
        detail = res.json().get("detail", "")
        assert "Servicio Ollama no disponible" in detail
        # Confirmar que NO se emitió el texto determinista con cifras de cama inventadas
        assert "Nairobi City registra 160 camas UCI" not in detail
        assert "RunnableLambda" not in detail
    finally:
        os.environ["OLLAMA_BASE_URL"] = orig_url


def test_provoke_failure_corrupt_simulation_json_rejects_without_fallback():
    """Provoca un registro de simulación corrupto o incompleto y confirma rechazo HTTP 422 sin generar PDF/Excel con ceros."""
    from fastapi.testclient import TestClient
    from backend.app.main import app
    from backend.app.db.session import SessionLocal
    from backend.app.models.simulation import SimulationRunDB
    from datetime import datetime

    client = TestClient(app)
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    db = SessionLocal()
    test_sim_id = "sim-corrupt-test-999"
    try:
        # 1. Insertar simulación con parameters_json corrupto (no JSON)
        corrupt_sim = SimulationRunDB(
            id=test_sim_id,
            user_id="test_admin_user",
            name="Simulación Corrupta de Prueba",
            run_mode="SEIR_METAPOPULATION",
            parameters_json="ESTO_NO_ES_UN_JSON_VALIDO{{{",
            results_summary_json="{}",
            created_at=datetime.utcnow()
        )
        db.add(corrupt_sim)
        db.commit()

        # Preview debe retornar 422 explícito
        prev_res = client.post(
            "/api/v1/reports/preview",
            json={"report_type": "scenarios", "simulation_id": test_sim_id},
            headers=headers
        )
        assert prev_res.status_code == 422, f"Se esperaba 422 para simulación corrupta, recibido {prev_res.status_code}"
        assert "corrupto" in prev_res.json().get("detail", "").lower()

        # Export PDF debe fallar con 422 y NO generar PDF binario
        pdf_res = client.post(
            "/api/v1/reports/export/pdf",
            json={"report_type": "scenarios", "simulation_id": test_sim_id},
            headers=headers
        )
        assert pdf_res.status_code == 422
        assert pdf_res.headers.get("content-type") == "application/json"

        # Export Excel debe fallar con 422 y NO generar Excel binario
        excel_res = client.post(
            "/api/v1/reports/export/excel",
            json={"report_type": "scenarios", "simulation_id": test_sim_id},
            headers=headers
        )
        assert excel_res.status_code == 422
        assert excel_res.headers.get("content-type") == "application/json"

        # 2. Probar con parameters válido pero incompleto (falta days)
        corrupt_sim.parameters_json = json.dumps({"npi_stringency": 0.5, "vaccine_coverage": 0.3, "seed": 42})
        db.commit()

        prev_res2 = client.post(
            "/api/v1/reports/preview",
            json={"report_type": "scenarios", "simulation_id": test_sim_id},
            headers=headers
        )
        assert prev_res2.status_code == 422
        assert "incompleto" in prev_res2.json().get("detail", "").lower()
        assert "days" in prev_res2.json().get("detail", "")
    finally:
        sim_to_del = db.query(SimulationRunDB).filter(SimulationRunDB.id == test_sim_id).first()
        if sim_to_del:
            db.delete(sim_to_del)
            db.commit()
        db.close()


